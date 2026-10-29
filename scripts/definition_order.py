"""Verify a permutation-only update without rerunning definition calculations."""
from __future__ import annotations
import hashlib
import html
import json
import re
from pathlib import Path
from html.parser import HTMLParser
from collections import Counter
from copy import deepcopy
import zipfile
from lxml import etree as ET

RECEIPT = 'definition_order_update.json'
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'

def digest(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda:f.read(1048576),b''):h.update(block)
    return h.hexdigest().upper()

class Node:
    def __init__(self,tag,attrs,start,head,parent=None):
        self.tag,self.attrs,self.start,self.head,self.parent=tag,dict(attrs),start,head,parent
        self.end=self.tail=None;self.children=[]
    def find(self,predicate):
        for n in self.children:
            if predicate(n):yield n
            yield from n.find(predicate)
    def ancestor(self,predicate):
        p=self.parent
        while p:
            if predicate(p):return p
            p=p.parent

class Spans(HTMLParser):
    def __init__(self,text):
        super().__init__(convert_charrefs=False)
        self.lines=[0]
        for m in re.finditer('\n',text):self.lines.append(m.end())
        self.root=Node('root',[],0,0);self.stack=[self.root];self.text=text;self.feed(text)
    def position(self):
        row,col=self.getpos();return self.lines[row-1]+col
    def handle_starttag(self,tag,attrs):
        s=self.position();n=Node(tag,attrs,s,s+len(self.get_starttag_text()),self.stack[-1]);self.stack[-1].children.append(n)
        if tag not in {'area','base','br','col','embed','hr','img','input','link','meta','param','source','track','wbr'}:self.stack.append(n)
        else:n.tail=n.head;n.end=n.head
    def handle_startendtag(self,tag,attrs):
        self.handle_starttag(tag,attrs)
        if self.stack[-1].tag==tag:self.stack.pop()
    def handle_endtag(self,tag):
        s=self.position()
        for i in range(len(self.stack)-1,0,-1):
            if self.stack[i].tag==tag:
                n=self.stack[i];n.tail=s;n.end=self.text.index('>',s)+1;del self.stack[i:];break

def patch(text,edits):
    edits=sorted(edits)
    if any(a[1]>b[0] for a,b in zip(edits,edits[1:])):raise ValueError('Overlapping order patches')
    for start,end,value in reversed(edits):text=text[:start]+value+text[end:]
    return text

def permute_slots(text,spans,keys,order):
    rank={v:i for i,v in enumerate(order)}
    if len(keys)!=len(set(keys)) or set(keys)!=set(order):raise ValueError('Order slots must contain the exact final set')
    values=[text[a:b] for a,b in spans]
    ordered=[values[i] for i in sorted(range(len(keys)),key=lambda i:rank[keys[i]])]
    return patch(text,[(a,b,v) for (a,b),v in zip(spans,ordered)])

def reorder_flat_tree(text,order,aliases,rendered=False):
    """Only flat parallel result leaves; never move a parent across its children."""
    if rendered:
        pattern=r'(<pre\b[^>]*data-summary-tree="true"[^>]*><code>)(.*?)(</code></pre>)'
        result=r'<span data-tree-role="result"(?: data-tree-variable="([^"]+)")?>([^<]+)</span>'
    else:
        pattern=r'(^```text[^\n]*\n)(.*?)(^```[ \t]*$)'
        result=r'\[\[result:(?:([A-Za-z][A-Za-z0-9_.]*)\|)?([^\]]+)\]\]'
    def change(m):
        lines=m[2].splitlines();marked=[(i,re.search(result,line)) for i,line in enumerate(lines)]
        nodes=[(i,match) for i,match in marked if match]
        if len(nodes)!=len(order) or not nodes:return m[0]
        start=nodes[0][0]
        if [i for i,_ in nodes]!=list(range(start,len(lines))):raise ValueError('Not a flat result tree')
        keys=[html.unescape(match[1]) if match[1] else aliases.get(html.unescape(match[2])) for _,match in nodes]
        if set(keys)!=set(order):raise ValueError('Flat tree result mapping differs')
        body={key:lines[i] for key,(i,_) in zip(keys,nodes)}
        rebuilt=[]
        for i,v in enumerate(order):
            line=body[v]
            if rendered:
                line=re.sub(r'(<span data-tree-role="connector">)[├└]─(</span>)',lambda n:n[1]+('└' if i==len(order)-1 else '├')+'─'+n[2],line,count=1)
            else:
                if not re.match(r'^[├└]─ ',line):raise ValueError('Tree has nested results')
                line=('└' if i==len(order)-1 else '├')+line[1:]
            rebuilt.append(line)
        return m[1]+'\n'.join(lines[:start]+rebuilt)+('\n' if not rendered else '')+m[3]
    return re.sub(pattern,change,text,flags=re.M|re.S)

def reorder_html(text,order,groups,flat_aliases=None):
    """Move complete existing rows/cards/groups, preserving their exact bytes."""
    group_order=list(dict.fromkeys(groups.get(v,'') for v in order))
    if group_order!=['']:
        # Overview business sections, including their source and defined roles.
        doc=Spans(text);edits=[]
        for table in doc.root.find(lambda n:n.tag=='table' and 'heatmap-table' in n.attrs.get('class','').split()):
            starts=[];last=None
            for row in table.find(lambda n:n.tag=='tr'):
                if row.ancestor(lambda n:n.tag=='table') is not table:continue
                marks=list(row.find(lambda n:n.attrs.get('data-component-title') or n.attrs.get('data-component-group')))
                if marks:
                    g=marks[0].attrs.get('data-component-title') or marks[0].attrs.get('data-component-group')
                    if g!=last:starts.append((row.start,g));last=g
            if starts:
                ends=[s for s,_ in starts[1:]]+[table.tail]
                chunks={g:text[s:e] for (s,g),e in zip(starts,ends)}
                if set(chunks)!=set(group_order):raise ValueError('Overview group set changed')
                edits.append((starts[0][0],table.tail,''.join(chunks[g] for g in group_order)))
        text=patch(text,edits)
        # Details use h3 sections in full mode, category details in compact mode.
        doc=Spans(text);edits=[]
        for container in doc.root.find(lambda n:n.tag=='div' and 'variable-map' in n.attrs.get('class','').split()):
            heads=[n for n in container.children if n.attrs.get('data-component-surface')=='details']
            if heads:
                spans=[(n.start,heads[i+1].start if i+1<len(heads) else container.tail) for i,n in enumerate(heads)]
                keys=[n.attrs['data-component-title'] for n in heads]
            else:
                cats=[n for n in container.children if n.tag=='details' and n.attrs.get('id','').startswith('cat-')]
                spans=[];keys=[]
                for n in cats:
                    mark=next(n.find(lambda e:e.tag=='summary' and e.attrs.get('data-component-group')),None)
                    if mark is not None:spans.append((n.start,n.end));keys.append(mark.attrs['data-component-group'])
            if spans:
                if len(keys)!=len(set(keys)) or set(keys)!=set(group_order):raise ValueError('Detail group set changed')
                values={k:text[a:b] for k,(a,b) in zip(keys,spans)}
                for (a,b),g in zip(spans,group_order):edits.append((a,b,values[g]))
        text=patch(text,edits)
    # Each result table has its own sequence; source rows stay in their slots.
    doc=Spans(text);edits=[]
    for table in doc.root.find(lambda n:n.tag=='table'):
        rows=[];keys=[]
        for row in table.find(lambda n:n.tag=='tr'):
            if row.ancestor(lambda n:n.tag=='table') is not table:continue
            raw=text[row.start:row.end]
            match=re.search(r'href=["\']#var-([^"\']+)',raw)
            variable=match[1] if match else None
            if variable not in order:
                cell=next(row.find(lambda n:n.tag=='td' and 'plain-cell' in n.attrs.get('class','').split()),None)
                variable=re.sub('<[^>]*>','',text[cell.head:cell.tail]).strip() if cell else None
            if variable in order:rows.append((row.start,row.end));keys.append(variable)
        if rows:
            reordered=permute_slots(text,rows,keys,order)
            for (a,b),key in zip(rows,order):
                index=keys.index(key);sa,sb=rows[index];edits.append((a,b,text[sa:sb]))
    text=patch(text,edits)
    doc=Spans(text);cards=[];keys=[]
    for n in doc.root.find(lambda n:n.tag=='div' and 'data-card' in n.attrs.get('class','').split()):
        if n.ancestor(lambda n:n.tag=='table'):continue
        mark=next(n.find(lambda x:x.attrs.get('id','').startswith('var-') and 'var-name' in x.attrs.get('class','').split()),None)
        if mark is not None and mark.attrs['id'][4:] in order:
            cards.append((n.start,n.end));keys.append(mark.attrs['id'][4:])
    if cards:text=permute_slots(text,cards,keys,order)
    return reorder_flat_tree(text,order,flat_aliases,True) if flat_aliases is not None else text

def reorder_copy(text,order,flat_aliases=None):
    from check_reader_copy import sections
    parts=sections(text,2)
    if flat_aliases is not None:text=reorder_flat_tree(text,order,flat_aliases)
    if 'Criteria' not in parts:return text
    blocks=sections(parts['Criteria'],3)
    if set(blocks)!=set(order):raise ValueError('Criteria variable set differs')
    # Keep headings and complete bodies byte-for-byte; change only block order.
    matches=list(re.finditer(r'(?m)^### ([^\r\n]+)[ \t\r]*$',parts['Criteria']))
    chunks={m[1].strip().strip('`'):parts['Criteria'][m.start():matches[i+1].start() if i+1<len(matches) else len(parts['Criteria'])].rstrip() for i,m in enumerate(matches)}
    intro=parts['Criteria'][:matches[0].start()]
    replacement=intro+'\n\n'.join(chunks[v] for v in order)+'\n\n'
    return re.sub(r'(?ms)(^## Criteria[^\n]*\n).*?(?=^## |\Z)',lambda m:m[1]+'\n'+replacement,text,count=1)

def r_call_arguments(text,name):
    start=text.index(name+'(')+len(name)+1;depth=0;quote=None;escape=False;comment=False;last=start;parts=[]
    for i in range(start,len(text)):
        c=text[i]
        if comment:
            if c=='\n':comment=False
        elif quote:
            if escape:escape=False
            elif c=='\\':escape=True
            elif c==quote:quote=None
        elif c in ['"',"'",'`']:quote=c
        elif c=='#':comment=True
        elif c in '([{':depth+=1
        elif c in ')]}':
            if c==')' and depth==0:parts.append((last,i));return parts
            depth-=1
        elif c==',' and depth==0:parts.append((last,i));last=i+1
    raise ValueError('Unclosed R call')

ORDER_LINES='''db_data <- db_data[c(setdiff(names(db_data), analysis_vars), analysis_vars)]
analysis_data <- analysis_data[c(setdiff(names(analysis_data), analysis_vars), analysis_vars)]
codebook <- codebook[match(names(db_data), codebook$Variable), , drop = FALSE]
analysis_codebook <- analysis_codebook[match(analysis_vars, analysis_codebook$Variable), , drop = FALSE]
'''

def reorder_r(text,order,groups):
    q=lambda x:json.dumps(x,ensure_ascii=False)
    match=re.search(r'(?ms)^analysis_vars\s*<-\s*c\((.*?)\)\s*$',text)
    if not match:raise ValueError('No static analysis_vars vector')
    before=re.findall(r'["\']([^"\']+)["\']',match[1])
    if set(before)!=set(order):raise ValueError('R variable set changed')
    if before!=order:
        text=text[:match.start()]+'analysis_vars <- c('+', '.join(map(q,order))+')'+text[match.end():]
        match=re.search(r'(?m)^.*(?:write\.xlsx|write\.csv)\s*\(',text)
        if not match:raise ValueError('No output write location')
        text=text[:match.start()]+ORDER_LINES+'\n'+text[match.start():]
    if groups and any(groups.values()):
        args=r_call_arguments(text,'render_definition_bundle');edits=[]
        gorder=list(dict.fromkeys(groups[v] for v in order))
        vector='c('+', '.join(q(v)+' = '+q(groups[v]) for v in order)+')'
        levels='c('+', '.join(map(q,gorder))+')'
        for a,b in args:
            content=text[a:b];m=re.match(r'(\s*)(summary_groups|component_groups)\s*=\s*(.*)',content,re.S)
            if not m:continue
            if m[2]=='summary_groups':value='factor('+vector+', levels = '+levels+')'
            else:
                if m[3].strip()=='NULL':continue
                value='factor('+m[3].strip()+', levels = '+levels+')'
            edits.append((a,b,m[1]+m[2]+' = '+value))
        text=patch(text,edits)
    return text

def column(n):
    out=''
    while n:n,r=divmod(n-1,26);out=chr(65+r)+out
    return out

def cell_key(cell):
    return (cell.get('t'),cell.get('s'),ET.tostring(cell.find(NS+'v')),ET.tostring(cell.find(NS+'is'))) if cell.find(NS+'v') is not None and cell.find(NS+'is') is not None else (
        cell.get('t'),cell.get('s'),cell.findtext(NS+'v'),''.join(cell.find(NS+'is').itertext()) if cell.find(NS+'is') is not None else None)

def worksheet(path):
    with zipfile.ZipFile(path) as z:
        sheets=[n for n in z.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',n)]
        if len(sheets)!=1:raise ValueError('Order-only update requires one worksheet')
        return sheets[0]

def workbook_header(path):
    with zipfile.ZipFile(path) as z,z.open(worksheet(path)) as f:
        first=next(raw for kind,raw in worksheet_chunks(f) if kind=='row')
        node=ET.fromstring(b'<root xmlns="'+NS[1:-1].encode()+b'">'+first+b'</root>')[0]
        cells=list(node);needed={int(c.findtext(NS+'v')) for c in cells if c.get('t')=='s'};strings={}
        if needed:
            with z.open('xl/sharedStrings.xml') as sf:
                index=0
                for _,n in ET.iterparse(sf,events=('end',),tag=NS+'si'):
                    if index in needed:strings[index]=''.join(n.itertext())
                    n.clear();index+=1
                    if len(strings)==len(needed):break
        return [strings[int(c.findtext(NS+'v'))] if c.get('t')=='s' else ''.join(c.itertext()) for c in cells]

def xlsx_permute(before,after,order,dictionary=False):
    """Keep every cell type/value/style; move only result columns or dictionary rows."""
    from openpyxl import load_workbook
    header=workbook_header(before)
    sheet=worksheet(before)
    if dictionary:
        b=load_workbook(before,read_only=True,data_only=False);b.active.reset_dimensions();entries=list(b.active.values);b.close()
        vi=header.index('Variable');names=[row[vi] for row in entries[1:]]
        slots=[i for i,n in enumerate(names) if n in order]
        if {names[i] for i in slots}!=set(order):raise ValueError('Dictionary variable set differs')
        target=[n for n in names if n not in order]+order
        with zipfile.ZipFile(before) as zin:
            root=ET.fromstring(zin.read(sheet));data=root.find(NS+'sheetData');rows=list(data)
            mapping={row[vi]:node for row,node in zip(entries[1:],rows[1:])}
            new=[rows[0]]+[mapping[n] for n in target]
            for ri,row in enumerate(new,1):
                row.set('r',str(ri))
                for cell in row:cell.set('r',re.sub(r'\d+$',str(ri),cell.get('r')))
            data[:]=new
            with zipfile.ZipFile(after,'w',zipfile.ZIP_DEFLATED,compresslevel=1) as zout:
                for info in zin.infolist():zout.writestr(info,ET.tostring(root,xml_declaration=True,encoding='UTF-8') if info.filename==sheet else zin.read(info.filename))
        return
    slots=[i for i,n in enumerate(header) if n in order]
    if {header[i] for i in slots}!=set(order):raise ValueError('Data result columns differ')
    target=[n for n in header if n not in order]+order
    permutation=[header.index(n) for n in target];rank={old+1:new+1 for new,old in enumerate(permutation)}
    with zipfile.ZipFile(before) as zin,zipfile.ZipFile(after,'w',zipfile.ZIP_DEFLATED,compresslevel=1) as zout:
        for info in zin.infolist():
            if info.filename!=sheet:zout.writestr(info,zin.read(info.filename));continue
            info._compresslevel=1
            with zin.open(sheet) as source,zout.open(info,'w',force_zip64=True) as dest:
                for kind,raw in worksheet_chunks(source):
                    if kind!='row':dest.write(raw);continue
                    if b'<f' in raw:raise ValueError('Formula cells require a formula-aware update')
                    matches=list(CELL.finditer(raw));cells=[]
                    for m in matches:
                        ref=REF.search(m[0]);letters=ref[1].decode();ci=0
                        for c in letters:ci=ci*26+ord(c)-64
                        ni=rank[ci]
                        value=m[0][:ref.start(1)]+column(ni).encode()+m[0][ref.end(1):]
                        cells.append((ni,value))
                    dest.write(raw[:matches[0].start()]+b''.join(v for _,v in sorted(cells))+raw[matches[-1].end():] if matches else raw)

CELL=re.compile(rb'<c\b[^>]*?(?:/>|>.*?</c>)',re.S)
REF=re.compile(rb'\br="([A-Z]+)([0-9]+)"')

def worksheet_chunks(source):
    buffer=b''
    while True:
        data=source.read(2*1024*1024)
        if not data:
            if buffer:yield 'other',buffer
            break
        buffer+=data
        cursor=0
        while True:
            start=buffer.find(b'<row ',cursor)
            if start<0:
                if len(buffer)-cursor>256:yield 'other',buffer[cursor:-256];buffer=buffer[-256:]
                else:buffer=buffer[cursor:]
                break
            end=buffer.find(b'</row>',start)
            if end<0:
                if start>cursor:yield 'other',buffer[cursor:start]
                buffer=buffer[start:]
                break
            if start>cursor:yield 'other',buffer[cursor:start]
            yield 'row',buffer[start:end+6]
            cursor=end+6

def xlsx_signature(path,dictionary=False):
    """Per-variable typed-cell hashes, independent of column/row positions."""
    from openpyxl import load_workbook
    header=workbook_header(path)
    if dictionary:
        b=load_workbook(path,read_only=True,data_only=False);b.active.reset_dimensions()
        rows=b.active.values;next(rows)
        index=header.index('Variable');data={}
        for row in rows:
            values=list(row);key=values[index]
            if key in data:raise ValueError('Duplicate dictionary variable')
            data[key]=values
        b.close();return dict(header=header,rows=data)
    canonical=sorted(header);indices=[header.index(v) for v in canonical]
    h=hashlib.sha256();count=0
    with zipfile.ZipFile(path) as z,z.open(worksheet(path)) as f:
        for kind,row in worksheet_chunks(f):
            if kind!='row':continue
            if count:
                cells={}
                for m in CELL.finditer(row):
                    ref=REF.search(m[0]);ci=0
                    for letter in ref[1]:ci=ci*26+letter-64
                    cells[ci-1]=m[0][:ref.start()]+m[0][ref.end():]
                # Entire typed cells, including styles, inline text and values.
                h.update(b'\x00'.join(cells.get(i,b'') for i in indices)+b'\n')
            count+=1
    return dict(rows=count-1,columns=canonical,typed_cells=h.hexdigest())

def verify(formal,process,topic):
    formal=Path(formal).resolve();process=Path(process).resolve()
    receipt=json.loads((process/RECEIPT).read_text(encoding='utf-8-sig'))
    if receipt['topic_id']!=str(topic).zfill(3) or receipt['formal']!=str(formal):raise ValueError('Order receipt identity differs')
    backup=Path(receipt['before']).resolve();backup.relative_to(process)
    order=receipt['order'];groups=receipt['groups']
    aliases=receipt.get('flat_aliases')
    if len(set(order))!=len(order):raise ValueError('Duplicate final variable')
    binding={'receipt':digest(process/RECEIPT),'verifier':digest(__file__),
      'current':{p.name:digest(p) for p in formal.iterdir() if p.is_file()},
      'original':{name:digest(backup/name) for name in receipt['changed_files']},
      'source':digest(process/'definition_search_record.json'),
      'original_source':digest(backup/'definition_search_record.json')}
    cache=backup.parent/'permutation_verified.json'
    if cache.is_file():
        saved=json.loads(cache.read_text(encoding='utf-8'))
        if saved.get('binding')==binding:return saved['result']
    for filename,sha in receipt['before_hashes'].items():
        p=backup/filename if filename in receipt['changed_files'] else formal/filename
        if digest(p)!=sha:raise ValueError('Order baseline changed: '+filename)
    if set(p.name for p in formal.iterdir() if p.is_file())!=set(receipt['before_hashes']):raise ValueError('Artifact set changed')
    for filename in receipt['changed_files']:
        before=backup/filename;after=formal/filename
        if after.suffix.lower()=='.xlsx':
            dictionary=filename.startswith(('codebook_','analysis_codebook_'))
            if xlsx_signature(before,dictionary)!=xlsx_signature(after,dictionary):raise ValueError('Values, names, labels or row sequence changed: '+filename)
        elif filename=='文案.md':
            expected=reorder_copy(before.read_text(encoding='utf-8-sig'),order,aliases)
            if expected!=after.read_text(encoding='utf-8-sig'):raise ValueError('Copy changed beyond Criteria order')
        elif after.suffix.lower()=='.r':
            if reorder_r(before.read_text(encoding='utf-8-sig'),order,groups)!=after.read_text(encoding='utf-8-sig'):raise ValueError('R changed beyond output order')
        elif after.suffix.lower() in ('.html','.md'):
            old=before.read_text(encoding='utf-8-sig')
            if after.suffix.lower()=='.md' and filename==receipt['artifacts']['note']:
                original_r=backup/receipt['artifacts']['public_r']
                if not original_r.exists():original_r=formal/receipt['artifacts']['public_r']
                old=replace_public_code(old,original_r,formal/receipt['artifacts']['public_r'])
            if reorder_html(old,order,groups,aliases)!=after.read_text(encoding='utf-8-sig'):raise ValueError('Display changed beyond order: '+filename)
        else:raise ValueError('Unsupported order artifact: '+filename)
    # Source-plan list changes are limited to permutations of existing entries.
    old=json.loads((backup/'definition_search_record.json').read_text(encoding='utf-8-sig'))
    new=json.loads((process/'definition_search_record.json').read_text(encoding='utf-8-sig'))
    expected=deepcopy(old);expected['approved_analysis_vars']=order
    if isinstance(expected.get('definition_plan'),list):
        entries={x.get('analysis_var',x.get('analysis_variable')):x for x in expected['definition_plan']}
        if set(entries)!=set(order):raise ValueError('Plan final set differs')
        expected['definition_plan']=[entries[v] for v in order]
    if expected!=new:raise ValueError('Source plan changed beyond order')
    from check_definition_output import DefinedSurfaceParser,parse_definition_table
    for filename in [receipt['artifacts']['note'],*receipt['detail_files']]:
        parser=DefinedSurfaceParser();parser.feed((formal/filename).read_text(encoding='utf-8-sig'))
        if [v for v in parser.overview_variables if v in order]!=order or [v for v,_ in parser.detail_groups if v in order]!=order:raise ValueError('Final display order differs: '+filename)
        if groups and any(groups.values()):
            if any(groups[v]!=g for v,g in parser.overview_groups+parser.detail_groups if v in order):raise ValueError('Business group membership changed')
    from openpyxl import load_workbook
    for filename in receipt['workbooks']:
        b=load_workbook(formal/filename,read_only=True,data_only=False);b.active.reset_dimensions();it=b.active.values;header=list(next(it))
        if filename.startswith(('codebook_','analysis_codebook_')):
            vi=header.index('Variable');values=[row[vi] for row in it if len(row)>vi]
        else:values=header
        b.close()
        if [v for v in values if v in order]!=order:raise ValueError('Workbook result order differs: '+filename)
    result={'ok':True,'status':'ORDER_UPDATE_READY','topic_id':str(topic).zfill(3),
      'scope':'order_only','writing_quality':'NOT_ASSESSED','naming_quality':'DEFERRED_BY_USER',
      'audit':str(process/RECEIPT),'audit_sha256':digest(process/RECEIPT),
      'change_impact':{'path':RECEIPT,'sha256':digest(process/RECEIPT),'changed_dimensions':['order']},
      'result_check':{'scope':'order_only','status':'PERMUTATION_CHECK_PASS'},
      'questionnaire_rendering':{'status':'PRESERVED_BY_EXACT_TRANSFORMATION'},
      'reader_review':{'status':'NOT_REPEATED_FOR_PERMUTATION'},
      'artifacts':{role:{'path':name,'sha256':digest(formal/name)} for role,name in receipt['artifacts'].items()},
      'validation_basis':'Existing content plus verified permutations; no recomputation or fresh business/naming review'}
    cache.write_text(json.dumps({'binding':binding,'result':result},ensure_ascii=False,indent=2),encoding='utf-8')
    return result

def replace_public_code(note,before,after):
    a=Path(before).read_text(encoding='utf-8-sig');b=Path(after).read_text(encoding='utf-8-sig')
    # The note embeds either the complete script or its public prefix.
    for old,new in [(a,b),(a.split('# 输出',1)[0].rstrip(),b.split('# 输出',1)[0].rstrip())]:
        if old in note:return note.replace(old,new,1)
    if a==b:return note
    raise ValueError('Public R cannot be uniquely located in note')
