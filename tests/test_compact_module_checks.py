"""Checker boundary and adversarial fixtures for explicit compact modules."""
from pathlib import Path
import html
import importlib.util
import json
import re
from collections import Counter
import sys
import tempfile
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from definition_display import resolve_display_policy
from test_definition_readability_gate import (complete_impact, load_checker as load_readability,
    initialize_fixture_audit, FIXTURE_CONTENT, render_note)


def load_output():
    spec = importlib.util.spec_from_file_location("compact_output", ROOT / "scripts/check_definition_output.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def surface(names):
    return ('<a href="#cat-defined">Defined variables</a>'
            '<summary>Defined variables</summary>'
            '<table data-defined-overview="true">' + ''.join(
                f'<tr><td><a href="#var-{html.escape(n, quote=True)}">{html.escape(n)}</a></td></tr>' for n in names)
            + '</table>' + ''.join(
                f'<span class="var-name" id="var-{html.escape(n, quote=True)}" data-defined-variable="{html.escape(n, quote=True)}">{html.escape(n)}</span>' for n in names))


def workbook(path, rows):
    def column(index):
        name = ""
        while index:
            index, remainder = divmod(index - 1, 26)
            name = chr(65 + remainder) + name
        return name

    with zipfile.ZipFile(path, "w") as z:
        z.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="data" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        body = ''.join(f'<row r="{i}">' + ''.join(f'<c r="{column(j)}{i}" t="inlineStr"><is><t>{html.escape(str(v))}</t></is></c>' for j, v in enumerate(row, 1)) + '</row>' for i, row in enumerate(rows, 1))
        z.writestr("xl/worksheets/sheet1.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + body + '</sheetData></worksheet>')


def assert_failure(fn):
    try:
        fn()
    except ValueError:
        return
    raise AssertionError("expected ValueError")


def main():
    c = load_output()
    from definition_order import xlsx_permute, xlsx_signature, reorder_html, reorder_r, reorder_flat_tree
    with tempfile.TemporaryDirectory() as tmp:
        from openpyxl import Workbook, load_workbook
        path=Path(tmp);book=Workbook();sheet=book.active
        sheet.append(['ID','B','A']);sheet.append(['001',3.25,None]);sheet.append(['002',0,'Never'])
        book.save(path/'before.xlsx');book.close()
        xlsx_permute(path/'before.xlsx',path/'after.xlsx',['A','B'])
        assert xlsx_signature(path/'before.xlsx')==xlsx_signature(path/'after.xlsx')
        from io import BytesIO
        actual=load_workbook(BytesIO((path/'after.xlsx').read_bytes()));assert list(actual.active.values)[0]==('ID','A','B')
        actual.active['B2']=4;actual.save(path/'changed.xlsx');actual.close()
        assert xlsx_signature(path/'before.xlsx')!=xlsx_signature(path/'changed.xlsx')
        markup='<table class="heatmap-table"><tr><th>Variable</th></tr><tr><td><a href="#var-B">B</a></td></tr><tr><td><a href="#var-A">A</a></td></tr></table><div class="data-card"><span class="var-name" id="var-B">B</span><p>unchanged B</p></div><div class="data-card"><span class="var-name" id="var-A">A</span><p>unchanged A</p></div>'
        ordered=reorder_html(markup,['A','B'],{})
        assert ordered.index('#var-A')<ordered.index('#var-B') and ordered.index('id="var-A"')<ordered.index('id="var-B"')
        assert 'unchanged A' in ordered and 'unchanged B' in ordered
        source='analysis_vars <- c("B", "A")\nwrite.xlsx(db_data,"out.xlsx")\nrender_definition_bundle(summary_groups=c(B="一",A="一"))'
        ordered=reorder_r(source,['A','B'],dict(A='一',B='一'))
        assert 'analysis_vars <- c("A", "B")' in ordered and 'write.xlsx(db_data,"out.xlsx")' in ordered
        flat='```text 结果\n结果\n├─ [[result:乙]] [[meta:（0/1）]]\n└─ [[result:甲]] [[meta:（0/1）]]\n```'
        changed=reorder_flat_tree(flat,['A','B'],{'甲':'A','乙':'B'})
        assert '├─ [[result:甲]]' in changed and '└─ [[result:乙]]' in changed
        assert Counter(re.findall(r'\[\[(?:result|meta):.*?\]\]',flat))==Counter(re.findall(r'\[\[(?:result|meta):.*?\]\]',changed))
        bound=flat.replace('result:乙','result:B|乙').replace('result:甲','result:A|甲')
        assert '├─ [[result:A|甲]]' in reorder_flat_tree(bound,['A','B'],{})
        rendered='<pre data-summary-tree="true"><code>结果\n<span data-tree-role="connector">├─</span> <span data-tree-role="result" data-tree-variable="B">乙</span>\n<span data-tree-role="connector">└─</span> <span data-tree-role="result" data-tree-variable="A">甲</span></code></pre>'
        reordered=reorder_flat_tree(rendered,['A','B'],{},True)
        assert reordered.index('data-tree-variable="A"') < reordered.index('data-tree-variable="B"')
    names = [f"V.{i}" for i in range(31)]
    for kind, count, compact in (("module", 10, False), ("module", 11, True), ("module", 30, True), ("questionnaire_module", 11, False), ("standard", 11, False)):
        copy = f"---\ntopic_kind: {kind}\n---\n"
        record = dict(topic_kind="module" if kind=="questionnaire_module" else kind, approved_analysis_vars=names[:count])
        assert resolve_display_policy(copy, record=record)["compact_module"] is compact
    assert not resolve_display_policy("legacy module with 100 results", names)["compact_module"]
    assert_failure(lambda: resolve_display_policy("---\ntopic_kind: module\n---\n"))
    assert_failure(lambda: resolve_display_policy("---\ntopic_kind: module\n---\n", names, dict(topic_kind="standard")))

    with tempfile.TemporaryDirectory() as tmp:
        formal = Path(tmp)
        detail = formal / "topic_detail.html"
        good = surface(names)

        def checks(note=good, detail_text=good):
            detail.write_text(detail_text, encoding="utf-8")
            results = []
            c.check_compact_surfaces(formal, note, names, results)
            c.check_category_order(formal, note, results, compact_module=True)
            return results

        assert all(r["ok"] for r in checks())
        # Business titles carry their role as metadata; never infer it from Chinese text.
        (formal / "文案.md").write_text("---\ntopic_kind: module\n---\n", encoding="utf-8")
        mapping = {n: "肝功能" if i < 15 else "肾功能" for i, n in enumerate(names)}
        def grouped_surface():
            overview, details = [], []
            for i, group in enumerate(dict.fromkeys(mapping.values())):
                subset = [n for n in names if mapping[n] == group]
                mark = f'data-component-group="{group}" data-variable-role="defined"'
                overview.append(f'<tr><td><a href="#cat-{i}" {mark}>{group}</a></td></tr>' + ''.join(
                    f'<tr><td><a href="#var-{n}">{n}</a></td></tr>' for n in subset))
                details.append(f'<summary {mark}>{group}</summary>' + ''.join(
                    f'<span class="var-name" id="var-{n}" data-defined-variable="{n}">{n}</span>' for n in subset))
            return '<table data-defined-overview="true">' + ''.join(overview) + '</table>' + ''.join(details)
        grouped = grouped_surface()
        planned = dict(topic_kind="module", approved_analysis_vars=names,
                       definition_plan=[dict(analysis_var=n, display_group=g) for n,g in mapping.items()])
        assert all(r["ok"] for r in checks(grouped,grouped))
        detail.write_text(grouped,encoding="utf-8")
        result=[];c.check_component_group_membership(formal,grouped,names,result,planned)
        assert result and all(r["ok"] for r in result),result
        for broken in (grouped.replace('>肝功能</summary>','>错误标题</summary>'),
                       grouped.replace('data-component-group="肝功能"','data-component-group="肾功能"',1),
                       grouped.replace('href="#var-V.0"','href="#var-V.30"')):
            result=[];c.check_component_group_membership(formal,broken,names,result,planned)
            assert any(not r['ok'] for r in result),result
        # Dots and escaped attribute values are identifiers, not CSS selectors.
        escaped = ['V.a&b', 'V."quoted"']
        parser = c.DefinedSurfaceParser()
        parser.feed(surface(escaped))
        assert parser.variables == parser.overview_variables == escaped
        for broken in (surface(names[:-1]), surface(names + [names[0]]), surface(names + ['Extra']), surface(list(reversed(names))),
                       good.replace('href="#var-V.0"', 'href="#var-Other"'),
                       good + '<span class="var-name" id="var-raw">raw</span>',
                       good.replace('id="var-V.0"', 'id="var-Other"'),
                       good + '<summary>Source variables</summary>',
                       good + '<table><tr><th>Criteria</th></tr></table>'):
            assert any(not r["ok"] for r in checks(broken)), broken[:100]
            assert any(not r["ok"] for r in checks(detail_text=broken)), broken[:100]
        (formal / "topic_definition.html").write_text("stale", encoding="utf-8")
        assert any(not r["ok"] for r in checks())
        (formal / "topic_definition.html").unlink()

        # Full mode also compares actual result order, ignoring source rows.
        result = []
        (formal / "文案.md").write_text("## 摘要导读\n无需关系树。", encoding="utf-8")
        detail.write_text(surface(list(reversed(names))), encoding="utf-8")
        c.check_result_presentation(formal, surface(names), names, {}, result)
        assert any(not r['ok'] and r['check'] == 'topic_detail.html final result order' for r in result), result

        # Final validation uses explicit variable identity, not Label equality.
        two = names[:2]
        detail.write_text(surface(two), encoding='utf-8')
        tree = '## 摘要导读\n```text 检验\n' + '\n'.join('[[result:' + v + '|中文指标]]' for v in two) + '\n```'
        (formal / '文案.md').write_text(tree, encoding='utf-8')
        result = []
        plan = {'definition_plan': [{'analysis_var': v, 'display_name': '中文' + str(i)} for i,v in enumerate(two)]}
        c.check_result_presentation(formal, surface(two), two, dict.fromkeys(two, 'English label'), result, plan)
        assert all(r['ok'] for r in result), result

        copy = "---\ntopic_kind: module\n---\n"
        (formal / "文案.md").write_text(copy, encoding="utf-8")
        record = dict(topic_kind="module", approved_analysis_vars=names)
        note = ('## 摘要导读\n' + '这是对变量模块、适用人群、调查时期和结果用途的完整说明。' * 3
                + '\n## 定义的组分概览\n' + good)
        (formal / "note.md").write_text(note, encoding="utf-8")
        detail.write_text(good, encoding="utf-8")
        workbook(formal / "cb.xlsx", [["Variable"]] + [[n] for n in names])
        workbook(formal / "db.xlsx", [["SEQN", "YEAR"] + names])
        # Public R outline has its own tests; isolate this test to real workbook/display coverage.
        c.check_public_code_outline = lambda *args: None
        results = []
        c.check_summary_facts(formal, "db.xlsx", "cb.xlsx", results, "nhanes", record=record)
        assert all(r["ok"] for r in results), results
        assert not any(r["check"] == "definition table variables and order" for r in results)
        text_results = []
        c.check_required_user_text(formal, ["适用人群"], text_results, dict(compact_module=True))
        assert all(r["ok"] for r in text_results), text_results
        workbook(formal / "db.xlsx", [["SEQN", "YEAR"] + names[:-1]])
        results = []
        c.check_summary_facts(formal, "db.xlsx", "cb.xlsx", results, "nhanes", record=record)
        assert any(not r["ok"] and r["check"] == "definition variables in analysis_db" for r in results)
        for kind, subset in (("module", names[:10]), ("questionnaire_module", names), ("standard", names)):
            (formal / "文案.md").write_text(f"---\ntopic_kind: {kind}\n---\n", encoding="utf-8")
            workbook(formal / "cb.xlsx", [["Variable"]] + [[n] for n in subset])
            results = []
            c.check_summary_facts(formal, "db.xlsx", "cb.xlsx", results, "nhanes", record=dict(topic_kind="module" if kind=="questionnaire_module" else kind, approved_analysis_vars=subset))
            assert any(not r["ok"] and r["check"] == "definition table structure" for r in results)

    r = load_readability()
    with tempfile.TemporaryDirectory() as tmp:
        formal, process = Path(tmp) / "formal", Path(tmp) / "process"
        formal.mkdir()
        complete_impact(r, formal, process)
        original = (formal / "文案.md").read_text(encoding="utf-8")
        without_criteria = original[:original.index("## Criteria")] + original[original.index("## 小book提示"):]
        for kind, count, accepted in (("module", 10, False), ("module", 11, True), ("questionnaire_module", 11, False), ("standard", 11, False)):
            (formal / "文案.md").write_text(f"---\ntopic_kind: {kind}\n---\n" + without_criteria, encoding="utf-8")
            (process / r.SOURCE_RECORD_NAME).write_text(json.dumps(dict(topic_kind="module" if kind=="questionnaire_module" else kind, approved_analysis_vars=names[:count])), encoding="utf-8")
            if accepted:
                r.validate_impact(formal, process, "025")
            else:
                try:
                    r.validate_impact(formal, process, "025")
                except ValueError:
                    pass
                else:
                    raise AssertionError("ordinary display must still require Criteria")

    # Use the same isolated artifact/result fixture as mechanical_delivery,
    # but actually bind compact copy, its source plan and 31-result workbook.
    from writing_evidence import scope_bindings
    with tempfile.TemporaryDirectory() as tmp:
        formal, process = Path(tmp) / "formal", Path(tmp) / "process"
        formal.mkdir()
        complete_impact(r, formal, process)
        copy_path = formal / "文案.md"
        original = copy_path.read_text(encoding="utf-8")
        compact_copy = ("---\ntopic_kind: module\n---\n"
                        + original[:original.index("## Criteria")]
                        + original[original.index("## 小book提示"):])
        copy_path.write_text(compact_copy, encoding="utf-8")
        r.write_json(process / r.SOURCE_RECORD_NAME, dict(schema_version=1,
            topic_kind="module", approved_analysis_vars=names))
        content = dict(FIXTURE_CONTENT, criteria={})
        note = render_note(content)
        note = note[:note.index("## 定义\n")] + "## 定义的组分概览\n" + surface(names) + "\n" + note[note.index("## 参考资料说明"):]
        files = dict(note="note.md", public_r="define.R", analysis_db="db.xlsx", analysis_codebook="cb.xlsx")
        (formal / files["note"]).write_text(note, encoding="utf-8")
        (formal / files["public_r"]).write_text("# isolated gate fixture\n", encoding="utf-8")
        workbook(formal / files["analysis_codebook"], [["Variable"]] + [[n] for n in names])
        workbook(formal / files["analysis_db"], [["SEQN", "YEAR"] + names, ["1", "2011"] + [0] * len(names)])
        initialize_fixture_audit(r, formal, process, "025", files)
        initialized = r.initialize_audit(formal, process, "025", files, overwrite=True)
        bindings = scope_bindings(formal, process, initialized["artifacts"])
        assert set(bindings) == {"title_summary", "definition_logic", "criteria", "insight_card", "references"}
        assert all(isinstance(digest, str) and len(digest) == 64 for digest in bindings.values())
        checked = r.validate_audit(formal, process, "025")
        assert checked["writing_quality"] == "NOT_ASSESSED" and checked["scope_count"] == 0
        assert r.verify_existing_readiness(formal, process, "025")["ok"]
        assert not list(formal.glob("*_definition.html"))
    print("COMPACT_MODULE_CHECKS_PASS: module10/11, questionnaire11 and standard11 boundaries, exact coverage, stale/source/duplicate rejection, real workbooks, scope_bindings and compact init/check/verify-ready")


if __name__ == "__main__":
    main()
